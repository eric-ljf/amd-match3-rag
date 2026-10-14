"""DOCX ingestion via python-docx. Tables become pipe-separated rows."""
import sys

from docx import Document as DocxDocument

from .base import Document


def ingest_docx(abs_path, rel_path):
    try:
        doc = DocxDocument(abs_path)
    except Exception as e:
        print(f"[ingest] skip unreadable docx {rel_path}: {e}", file=sys.stderr)
        return []
    parts = []
    for para in doc.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for ti, table in enumerate(doc.tables):
        parts.append(f"[Table {ti + 1}]")
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = "\n".join(parts)
    if not text.strip():
        return []
    return [Document(file=rel_path, locator="document", text=text)]
